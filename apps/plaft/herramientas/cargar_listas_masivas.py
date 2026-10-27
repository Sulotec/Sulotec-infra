#!/usr/bin/env python3
"""Carga masiva de listas (PEP, sanciones, investigados, noticias, fraudes) en la base de PLAFT.

Lee los archivos que le des y reconoce solos su formato:
  - la PLANTILLA oficial (hoja DATOS, .xlsx o .csv)                 -> ver apps/plaft/plantillas/
  - «Base total.xlsx» de ASBANC (hoja «Base total» o «ASBANC»)        -> se convierte a la plantilla al vuelo
  - la hoja RISK de «base antigua pep.xlsx» (PEP nacional con cargo)  -> se convierte a la plantilla al vuelo

Es SEGURA por defecto: sin --aplicar solo SIMULA (lee la base pero no escribe nada) y muestra el informe.
Es repetible: no duplica (busca a la persona por tipo + numero de documento) y NUNCA borra ni cambia lo que ya existe.
Las filas con error no se cargan y se explican en un archivo aparte (--rechazados). Este programa no imprime datos personales.

Uso:
  python cargar_listas_masivas.py --archivo "Base total.xlsx" --archivo "base antigua pep.xlsx"            # simula
  python cargar_listas_masivas.py --archivo "Base total.xlsx" --archivo "base antigua pep.xlsx" --aplicar   # escribe
La conexion sale de DB_URL, DB_USER y DB_PASSWORD (variables de entorno), igual que cargar_datos.py.
"""
import argparse
import csv
import hashlib
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from datetime import date, datetime

import openpyxl

try:
    import pymysql
except ImportError:      # solo hace falta cuando se usa la base de datos
    pymysql = None

HOY = date.today()
FECHA_MIN = date(1930, 1, 1)
MIN_NACIMIENTO = date(1900, 1, 1)
LOTE = 4000
GRUPOS = {  # orden: (nombre, color) - los mismos de cargar_datos.py
    1: ('Listas Restrictivas', 'Vinotinto'),
    2: ('Listas Asociadas a LA/FT o Corrupcion (Penal)', 'Rojo'),
    3: ('Listas Asociadas a LA/FT o Corrupcion (Administrativo)', 'Naranja'),
    4: ('Sanciones Administrativas', 'Amarillo'),
    5: ('Listas de Afectacion Financiera', 'Amarillo'),
    6: ('Listas Informativas y PEPs', 'Verde'),
}
TIPOS_DOC = {
    'DNI': 'Documento Nacional de Identidad',
    'RUC': 'Registro Unico de Contribuyentes',
    'CE': 'Carne de Extranjeria',
    'PASAPORTE': 'Pasaporte',
    'SIN DOCUMENTO': 'Sin documento registrado',
}
# Listas a las que se manda cada cosa. clave interna: (nombre si hay que crearla, regex para encontrarla en el catalogo, grupo, alcance, es_pep, descripcion)
LISTAS = {
    'PEP_PERU': ('Perú - Personas Políticamente Expuestas (Resolución SBS N° 4349-2016)', r'^PERU - PERSONAS POLITICAMENTE EXPUESTAS \(RESOLUCION SBS', 6, 'NACIONAL', 1,
                 'PEP peruanos'),
    'PEP_GENERICA': ('Personas Expuestas Públicamente', r'^PERSONAS EXPUESTAS PUBLICAMENTE$', 6, 'NACIONAL E INTERNACIONAL', 1, 'PEP nacionales e internacionales'),
    'CONTRALORIA': ('Perú - Registro de Sanciones por Responsabilidad Administrativa Funcional a cargo de la Contraloría General de la República',
                    r'RESPONSABILIDAD ADMINISTRATIVA FUNCIONAL', 3, 'NACIONAL', 0, 'Sanciones de la Contraloria'),
    'OFAC': ('Lista OFAC', r'^LISTA OFAC$', 1, 'INTERNACIONAL', 0, 'Lista de sanciones de la OFAC'),
    'ONU': ('Internacional - Consejo de Seguridad ONU', r'^INTERNACIONAL - CONSEJO DE SEGURIDAD ONU$', 1, 'INTERNACIONAL', 0, 'Lista del Consejo de Seguridad de la ONU'),
    'SBS_SANC': ('Perú - Personas sancionadas por la Superintendencia de Banca y Seguros', r'^PERU - PERSONAS SANCIONADAS POR LA SUPERINTENDENCIA DE BANCA Y SEGUROS$',
                 4, 'NACIONAL', 0, 'Sujetos obligados sancionados por la SBS'),
    'MINJUS': ('Perú - Deudores de reparaciones civiles (MINJUS)', r'DEUDORES DE REPARACIONES CIVILES', 2, 'NACIONAL', 0, 'Registro de deudores de reparaciones civiles del MINJUS (ASBANC)'),
    'UIF_SENT': ('Perú - Sentenciados por corrupción (informado por la UIF)', r'SENTENCIADOS POR CORRUPCION', 2, 'NACIONAL', 0, 'Sentenciados por corrupcion informados por la UIF (ASBANC)'),
    'ASB_INVEST': ('ASBANC - Investigados', r'^ASBANC - INVESTIGADOS$', 2, 'NACIONAL E INTERNACIONAL', 0, 'Personas investigadas (base ASBANC, tipo de lista 1)'),
    'ASB_WORLD': ('ASBANC - World-Check', r'^ASBANC - WORLD-CHECK$', 3, 'INTERNACIONAL', 0, 'Perfiles World-Check (base ASBANC, tipo de lista 1)'),
    'ASB_NOTICIAS': ('ASBANC - Noticias (información negativa)', r'^ASBANC - NOTICIAS', 6, 'NACIONAL E INTERNACIONAL', 0, 'Noticias de lavado, corrupcion, narcotrafico y terrorismo (ASBANC)'),
    'ASB_FRAUDES': ('ASBANC - Fraudes (información negativa)', r'^ASBANC - FRAUDES', 3, 'NACIONAL', 0, 'Informacion negativa por fraude (ASBANC)'),
    'ASB_SANC_INT': ('ASBANC - Sanciones internacionales (otras fuentes)', r'^ASBANC - SANCIONES INTERNACIONALES', 1, 'INTERNACIONAL', 0, 'Sanciones internacionales de otras fuentes (ASBANC)'),
    'ASB_OTRAS': ('ASBANC - Otras listas (tipo 2)', r'^ASBANC - OTRAS LISTAS \(TIPO 2\)$', 2, 'NACIONAL E INTERNACIONAL', 0, 'Otras listas del tipo 2 de ASBANC'),
    'RISK_PEP': ('Perú - PEP base antigua (RISK)', r'^PERU - PEP BASE ANTIGUA \(RISK\)$', 6, 'NACIONAL', 1, 'PEP nacionales de la base antigua RISK (funcionarios, fundadores y candidatos de organizaciones politicas)'),
}


# ----------------------------------------------------------------------------- utilidades
def sin_tildes(t):
    return unicodedata.normalize('NFD', str(t)).encode('ascii', 'ignore').decode()


def clave(t):
    return re.sub(r'\s+', ' ', sin_tildes(t)).strip().upper()


def limpio(v, largo=None):
    """Texto sin saltos ni espacios de mas; vacio -> None. Los numeros guardados como decimal (123.0) vuelven a entero."""
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    t = re.sub(r'[\x00-\x1f\x7f]', ' ', str(v))
    t = re.sub(r'\s+', ' ', t).strip().strip('"').strip()
    if not t or t.upper() in ('NULL', 'NAN', 'N/A', '-', '--'):
        return None
    return t[:largo] if largo else t


def a_fecha(v, minimo=FECHA_MIN):
    """Fecha valida o None (00/00/00, «/ /», NO INDICA y fechas imposibles dan None)."""
    d = None
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    else:
        t = limpio(v)
        if not t:
            return None
        m = re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?', t)
        if m:
            y, mo, dd = int(m.group(1)), int(m.group(2)), int(m.group(3))
        else:
            m = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', t)
            if not m:
                return None
            dd, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            d = date(y, mo, dd)
        except ValueError:
            return None
    return d if minimo <= d <= date(2100, 12, 31) else None


def cod_hash(*partes):
    return hashlib.sha1('|'.join(clave(p or '') for p in partes).encode()).hexdigest()[:12].upper()


def cargar_paises():
    """Codigo numerico ISO (604) -> nombre en espanol, desde paises_iso.csv (junto a este archivo)."""
    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'paises_iso.csv')
    mapa = {}
    if os.path.exists(ruta):
        with open(ruta, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f):
                mapa[r['numerico'].lstrip('0') or '0'] = (r['alfa2'], r['nombre_es'])
    return mapa


PAISES_ISO = cargar_paises()


def pais_de(valor):
    """Codigo numerico o de 2 letras -> nombre del pais; 999 / desconocido -> None."""
    t = limpio(valor)
    if not t:
        return None
    if t.isdigit():
        r = PAISES_ISO.get(t.lstrip('0') or '0')
        return None if not r or r[0] == 'ZZ' else r[1]
    if len(t) == 2:
        for a2, nombre in PAISES_ISO.values():
            if a2 == t.upper():
                return None if a2 == 'ZZ' else nombre
    return t[:100]


# ----------------------------------------------------------------------------- documentos
def doc_plantilla(tipo, numero):
    """Validacion ESTRICTA de la plantilla. Devuelve (tipo, documento, error)."""
    n = re.sub(r'\s+', '', numero or '')
    if tipo == 'DNI':
        ok = re.fullmatch(r'\d{8}', n)
    elif tipo == 'RUC':
        ok = re.fullmatch(r'\d{11}', n)
    elif tipo == 'CE':
        ok = re.fullmatch(r'[A-Za-z0-9]{6,12}', n)
    elif tipo == 'PASAPORTE':
        ok = re.fullmatch(r'[A-Za-z0-9]{5,12}', n)
    else:
        return 'SIN DOCUMENTO', (n[:50] or None), None
    return (tipo, n.upper(), None) if ok else (tipo, None, f'el número no tiene el formato de {tipo}')


def doc_asbanc(codigo, numero):
    """Codigos de ASBANC: 21 = DNI, 9 = RUC, 5 = pasaporte. El 35 y los demas son codigos internos: no son documentos reales."""
    cod = str(codigo).strip()
    n = re.sub(r'\s+', '', limpio(numero) or '')
    base = re.sub(r'[^0-9A-Za-z]', '', n)
    if cod == '21' and re.fullmatch(r'\d{5,8}', base) and re.fullmatch(r'[0-9?]+', n):
        return 'DNI', base.zfill(8)
    if cod == '9' and re.fullmatch(r'\d{11}', base) and re.fullmatch(r'[0-9?]+', n):
        return 'RUC', base
    if cod == '5' and re.fullmatch(r'[A-Za-z0-9]{5,12}', base):
        return 'PASAPORTE', base.upper()
    return 'SIN DOCUMENTO', f'ASBANC-{cod}-{n}'.upper()[:50]


# ----------------------------------------------------------------------------- reglas de listas de ASBANC
def listas_asbanc(tipo, obs):
    o = clave(obs or '')
    if tipo == '3':
        return ['PEP_PERU'] + (['CONTRALORIA'] if 'INHABILITADO POR LA CONTRALORIA' in o else [])
    if tipo in ('5', '2'):
        if 'OFAC' in o:
            return ['OFAC']
        if re.search(r'NACIONES UNIDAS|UNITED NATIONS|SECURITY COUNCIL|CONSEJO DE SEGURIDAD|\bONU\b', o):
            return ['ONU']
        if tipo == '2':
            if 'MINJUS' in o or 'REPARAC' in o:
                return ['MINJUS']
            if re.search(r'SENTENCIAD|SENTECIAD|CORRUPCION|\bUIF\b', o):
                return ['UIF_SENT']
            return ['ASB_OTRAS']
        return ['ASB_SANC_INT']
    if tipo == '1':
        if 'POLITICAL' in o:
            return ['PEP_GENERICA']
        if re.search(r'WORLD ?-?CHECK', o):
            return ['ASB_WORLD']
        return ['ASB_INVEST']
    if tipo == '6':
        return ['ASB_FRAUDES']
    if tipo == '7':
        return ['SBS_SANC'] if 'SO SANCIONADO' in o else ['ASB_NOTICIAS']
    return []


# ----------------------------------------------------------------------------- acumulador
class Acum:
    def __init__(self):
        self.ent = {}          # (tipo_doc, documento) -> (tipo_entidad, pais, ape_pat, ape_mat, nombres, razon_social, alias, sexo, fecha_nac)
        self.manchas = []      # (clave_entidad, lista, descripcion, link, f_desde, f_hasta, institucion, cargo, tipo_pep, p_desde, p_hasta)
        self.vistos = set()
        self.rechazos = Counter()
        self.avisos = Counter()
        self.filas = Counter()
        self.por_lista = Counter()
        self.por_tipo_doc = Counter()
        self.ejemplos = {}

    def agregar(self, tipo_doc, documento, ent, lista, descripcion, link, f_desde, f_hasta, inst, cargo, tipo_pep, p_desde, p_hasta):
        k = (tipo_doc, documento)
        if k not in self.ent:
            self.ent[k] = ent
            self.por_tipo_doc[tipo_doc] += 1
        h = hash((k, lista, descripcion, cargo, inst, f_desde))
        if h in self.vistos:
            self.avisos['aparición repetida en el archivo (se carga una sola vez)'] += 1
            return
        self.vistos.add(h)
        self.manchas.append((k, lista, descripcion, link, f_desde, f_hasta, inst, cargo, tipo_pep, p_desde, p_hasta))
        self.por_lista[lista] += 1


class Archivo:
    """Contexto de lectura de una hoja: cuenta filas y registra los rechazos sin guardar datos personales en pantalla."""

    def __init__(self, acum, nombre, hoja, rechazados):
        self.acum, self.nombre, self.hoja, self.salida = acum, nombre, hoja, rechazados
        self.leidas = 0

    def rechazar(self, nro, motivo, fila):
        self.acum.rechazos[motivo] += 1
        self.acum.ejemplos.setdefault(motivo, []).append(nro) if len(self.acum.ejemplos.get(motivo, [])) < 5 else None
        if self.salida:
            self.salida.writerow([self.nombre, self.hoja, nro, motivo] + ['' if c is None else str(c)[:200] for c in fila])


def indices(cabecera, requeridas):
    """Busca las columnas por su encabezado (sin tildes; en estos Excel una tilde viene como «?»)."""
    simples = [re.sub(r'[^A-Z0-9]+', ' ', clave(c) if c is not None else '').strip() for c in cabecera]
    ix, faltan = {}, []
    for nombre, patron in requeridas.items():
        pos = next((i for i, h in enumerate(simples) if re.fullmatch(patron, h)), None)
        if pos is None:
            faltan.append(nombre)
        ix[nombre] = pos
    return ix, faltan


def filas_de(ws):
    cab = None
    for nro, f in enumerate(ws.iter_rows(values_only=True), start=1):
        if cab is None:
            if any(c not in (None, '') for c in f):
                cab = f
                yield nro, None, cab
            continue
        if any(c not in (None, '') for c in f):
            yield nro, f, cab


# ----------------------------------------------------------------------------- lectores
PLACEHOLDER_FECHA = re.compile(r'[/\s.\-]*|0+/0+/0+|NO INDICA', re.I)


def leer_asbanc(ws, ctx):
    req = {'tipo_lista': r'TIPO DE LISTA', 'tipo_doc': r'TIPO DE DOCUMENTO', 'doc': r'N\w? ?MERO DE DOCUMENTO|NUMERO DE DOCUMENTO', 'pais': r'PA\w? ?S|PAIS',
           'denom': r'DENOMINACI\w? ?N|DENOMINACION', 'obs': r'OBS DE INHABILITACI\w? ?N|OBS DE INHABILITACION', 'desde': r'FECHA DESDE',
           'hasta': r'FECHA HASTA', 'ape2': r'SEGUNDO APELLIDO', 'ape1': r'APELLIDO', 'nom2': r'SEGUNDO NOMBRE', 'nom1': r'NOMBRE'}
    ix = None
    for nro, f, cab in filas_de(ws):
        if f is None:
            ix, faltan = indices(cab, req)
            if faltan:
                sys.exit(f'En «{ctx.hoja}» faltan las columnas: {", ".join(faltan)}')
            continue
        ctx.leidas += 1
        g = lambda k: f[ix[k]] if ix[k] is not None and ix[k] < len(f) else None
        tl = limpio(g('tipo_lista'))
        if not tl or not tl.isdigit():
            ctx.rechazar(nro, 'fila corrida: el tipo de lista no es un número', f)
            continue
        raw_desde = g('desde')
        desde = a_fecha(raw_desde)
        if desde is None and isinstance(raw_desde, str) and not PLACEHOLDER_FECHA.fullmatch(raw_desde.strip()):
            ctx.rechazar(nro, 'fila corrida: la fecha desde no es una fecha', f)
            continue
        ape_pat, ape_mat = limpio(g('ape1'), 255), limpio(g('ape2'), 255)
        nombres = limpio(' '.join(x for x in (limpio(g('nom1')), limpio(g('nom2'))) if x), 255)
        denom = limpio(g('denom'), 255)
        if ape_pat or nombres:
            tipo_ent, razon = 'NATURAL', None
        elif denom:
            tipo_ent, razon = 'JURIDICA', denom
        else:
            ctx.rechazar(nro, 'sin nombre ni denominación', f)
            continue
        numero = limpio(g('doc'))
        tipo_doc, documento = doc_asbanc(g('tipo_doc'), numero) if numero else ('SIN DOCUMENTO', 'ASBANC-SD-' + cod_hash(ape_pat, ape_mat, nombres, razon))
        obs = limpio(g('obs'), 2000)
        destinos = listas_asbanc(tl, obs)
        if not destinos:
            ctx.rechazar(nro, f'tipo de lista {tl} desconocido', f)
            continue
        ent = (tipo_ent, pais_de(g('pais')), ape_pat, ape_mat, nombres, razon, None, None, None)
        hasta = a_fecha(g('hasta'))
        for lista in destinos:
            ctx.acum.agregar(tipo_doc, documento, ent, lista, obs, None, desde, hasta, None, None, None, None, None)


def periodo(v):
    """(texto corto original, fecha o None) de un periodo del PEP; NO INDICA queda como texto y sin fecha."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    if isinstance(v, (datetime, date)):
        d = a_fecha(v)
        return (d.isoformat() if d else None), d
    t = limpio(v, 50)
    return t, a_fecha(t)


def leer_risk(ws, ctx):
    req = {'tipo': r'TIPO', 'doi': r'NUMERO DOI', 'pat': r'APELLIDO PATERNO', 'mat': r'APELLIDO MATERNO', 'nombres': r'NOMBRES', 'entidad': r'ENTIDAD',
           'cargo': r'CARGO', 'tipo_pep': r'TIPO PEP', 'desde': r'PERIODO DESDE', 'hasta': r'PERIODO HASTA'}
    ix = None
    for nro, f, cab in filas_de(ws):
        if f is None:
            ix, faltan = indices(cab, req)
            if faltan:
                sys.exit(f'En «{ctx.hoja}» faltan las columnas: {", ".join(faltan)}')
            continue
        ctx.leidas += 1
        g = lambda k: f[ix[k]] if ix[k] is not None and ix[k] < len(f) else None
        ape_pat = limpio(g('pat'), 255)
        ape_mat = limpio(g('mat'), 255)
        if ape_mat and ape_mat.isdigit():
            ape_mat = None
        nombres = limpio(g('nombres'), 255)
        if not (ape_pat or nombres):
            ctx.rechazar(nro, 'sin nombre ni apellido', f)
            continue
        doi = re.sub(r'\s+', '', limpio(g('doi')) or '')
        if re.fullmatch(r'\d{8}', doi) and clave(g('tipo') or 'DNI') in ('DNI', ''):
            tipo_doc, documento = 'DNI', doi
        elif doi:
            tipo_doc, documento = 'SIN DOCUMENTO', f'RISK-{doi}'.upper()[:50]
            ctx.acum.avisos['documento de RISK que no es un DNI de 8 dígitos (se guarda como sin documento)'] += 1
        else:
            tipo_doc, documento = 'SIN DOCUMENTO', 'SIN-DOC-' + cod_hash(ape_pat, ape_mat, nombres)
            ctx.acum.avisos['persona de RISK sin documento (se guarda como sin documento)'] += 1
        p_desde, f_desde = periodo(g('desde'))
        p_hasta, f_hasta = periodo(g('hasta'))
        cargo, inst = limpio(g('cargo'), 255), limpio(g('entidad'), 255)
        ent = ('NATURAL', 'Peru', ape_pat, ape_mat, nombres, None, None, None, None)
        ctx.acum.agregar(tipo_doc, documento, ent, 'RISK_PEP', 'Base antigua RISK (PEP nacional)', None, f_desde, f_hasta, inst, cargo,
                         limpio(g('tipo_pep'), 50) or 'NACIONAL', p_desde, p_hasta)


def clave_lista(nombre):
    """Si la lista de la plantilla es una de las que ya conoce este programa, usa la misma clave (así se cuentan juntas)."""
    k = clave(nombre)
    for clv, v in LISTAS.items():
        if re.search(v[1], k):
            return clv
    return 'PLANTILLA:' + nombre


COLUMNAS_PLANTILLA = ['tipo_entidad', 'tipo_documento', 'numero_documento', 'pais', 'apellido_paterno', 'apellido_materno', 'nombres', 'razon_social', 'alias',
                      'sexo', 'fecha_nacimiento', 'lista', 'detalle', 'enlace', 'fecha_desde', 'fecha_hasta', 'institucion', 'cargo', 'tipo_pep']


def leer_plantilla(filas, ctx, catalogo):
    """filas: iterador de (nro, lista de valores) con el encabezado ya consumido en `cab`."""
    for nro, fila in filas:
        ctx.leidas += 1
        d = dict(zip(COLUMNAS_PLANTILLA, fila))
        g = lambda k: limpio(d.get(k))
        tipo_ent = (g('tipo_entidad') or '').upper()
        if tipo_ent not in ('NATURAL', 'JURIDICA'):
            ctx.rechazar(nro, 'tipo_entidad no es NATURAL ni JURIDICA', fila)
            continue
        tipo = (g('tipo_documento') or '').upper()
        if tipo not in TIPOS_DOC:
            ctx.rechazar(nro, 'tipo_documento no permitido', fila)
            continue
        ape_pat, ape_mat, nombres, razon = (limpio(d.get(k), 255) for k in ('apellido_paterno', 'apellido_materno', 'nombres', 'razon_social'))
        if tipo_ent == 'NATURAL' and not (ape_pat or nombres):
            ctx.rechazar(nro, 'persona sin apellido ni nombres', fila)
            continue
        if tipo_ent == 'JURIDICA' and not razon:
            ctx.rechazar(nro, 'empresa sin razón social', fila)
            continue
        tipo_doc, documento, error = doc_plantilla(tipo, g('numero_documento'))
        if error:
            ctx.rechazar(nro, error, fila)
            continue
        if tipo_doc == 'SIN DOCUMENTO' and not documento:
            documento = 'SIN-DOC-' + cod_hash(ape_pat, ape_mat, nombres, razon)
        lista = g('lista')
        if not lista:
            ctx.rechazar(nro, 'falta la lista', fila)
            continue
        if catalogo is not None and clave(lista) not in catalogo:
            ctx.rechazar(nro, 'la lista no existe en el catálogo', fila)
            continue
        for campo in ('fecha_nacimiento', 'fecha_desde', 'fecha_hasta'):
            if g(campo) and a_fecha(d.get(campo), MIN_NACIMIENTO) is None:
                ctx.acum.avisos[f'{campo} no es una fecha válida (se deja vacía)'] += 1
        ent = (tipo_ent, pais_de(d.get('pais')) or 'Peru', ape_pat, ape_mat, nombres, razon, limpio(d.get('alias'), 255),
               (g('sexo') or '')[:1].upper() or None, a_fecha(d.get('fecha_nacimiento'), MIN_NACIMIENTO))
        ctx.acum.agregar(tipo_doc, documento, ent, clave_lista(lista), limpio(d.get('detalle'), 2000), limpio(d.get('enlace')),
                         a_fecha(d.get('fecha_desde')), a_fecha(d.get('fecha_hasta')), limpio(d.get('institucion'), 255), limpio(d.get('cargo'), 255),
                         limpio(d.get('tipo_pep'), 50), None, None)


def procesar_archivo(ruta, acum, rechazados, catalogo, omitir_hojas):
    nombre = os.path.basename(ruta)
    print(f'\n== {nombre}')
    if ruta.lower().endswith('.csv'):
        with open(ruta, encoding='utf-8-sig', newline='') as f:
            muestra = f.read(4096)
            f.seek(0)
            dialecto = csv.Sniffer().sniff(muestra, delimiters=',;\t') if muestra else csv.excel
            filas = csv.reader(f, dialecto)
            cab = [clave(c).lower() for c in next(filas, [])]
            if 'tipo_entidad' not in cab:
                sys.exit(f'{nombre}: el CSV no tiene el encabezado de la plantilla (tipo_entidad, tipo_documento, ...)')
            orden = [cab.index(c) if c in cab else None for c in COLUMNAS_PLANTILLA]
            ctx = Archivo(acum, nombre, 'CSV', rechazados)
            leer_plantilla(((n, [(fila[i] if i is not None and i < len(fila) else None) for i in orden]) for n, fila in enumerate(filas, start=2)), ctx, catalogo)
            print(f'   CSV (plantilla): {ctx.leidas} filas')
            acum.filas[nombre] += ctx.leidas
        return
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title in omitir_hojas:
            print(f'   hoja «{ws.title}»: ' + ('es de ejemplos, no se carga' if ws.title.upper() == 'EJEMPLOS' else 'se omite porque lo pediste (ya está completa en otro archivo)'))
            continue
        ctx = Archivo(acum, nombre, ws.title, rechazados)
        primera = next(ws.iter_rows(min_row=1, max_row=30, values_only=True), None)
        cabeceras = []
        for fila in ws.iter_rows(min_row=1, max_row=30, values_only=True):
            if any(c not in (None, '') for c in fila):
                cabeceras = [re.sub(r'[^A-Z0-9]+', ' ', clave(c) if c is not None else '').strip() for c in fila]
                break
        t0 = time.time()
        if 'TIPO ENTIDAD' in cabeceras and 'LISTA' in cabeceras:
            orden = [cabeceras.index(re.sub(r'[^A-Z0-9]+', ' ', c.upper()).strip()) if re.sub(r'[^A-Z0-9]+', ' ', c.upper()).strip() in cabeceras else None
                     for c in COLUMNAS_PLANTILLA]
            filas = ((nro, [(f[i] if i is not None and i < len(f) else None) for i in orden]) for nro, f, _ in filas_de(ws) if f is not None)
            leer_plantilla(filas, ctx, catalogo)
            formato = 'plantilla'
        elif 'TIPO DE LISTA' in cabeceras:
            leer_asbanc(ws, ctx)
            formato = 'ASBANC'
        elif 'NUMERO DOI' in cabeceras:
            leer_risk(ws, ctx)
            formato = 'RISK (PEP)'
        else:
            if not cabeceras:
                print(f'   hoja «{ws.title}»: vacía, se omite')
            else:
                print(f'   hoja «{ws.title}»: formato no reconocido, se omite')
            continue
        acum.filas[f'{nombre}/{ws.title}'] = ctx.leidas
        print(f'   hoja «{ws.title}» ({formato}): {ctx.leidas} filas leídas en {time.time() - t0:.0f} s')


# ----------------------------------------------------------------------------- base de datos
def leer_env(ruta):
    datos = {}
    if ruta and os.path.exists(ruta):
        for linea in open(ruta, encoding='utf-8'):
            m = re.match(r'\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*?)\s*$', linea)
            if m and not linea.lstrip().startswith('#'):
                datos[m.group(1)] = m.group(2)
    return datos


def conectar(args):
    if pymysql is None:
        sys.exit('Falta la librería pymysql.')
    env = leer_env(args.env)
    url = args.db_url or os.environ.get('DB_URL') or env.get('DB_URL')
    usuario = args.usuario or os.environ.get('DB_USER') or env.get('DB_USER')
    clave_db = args.clave if args.clave is not None else os.environ.get('DB_PASSWORD', env.get('DB_PASSWORD'))
    if not (url and usuario and clave_db is not None):
        sys.exit('Faltan DB_URL, DB_USER y DB_PASSWORD.')
    m = re.match(r'jdbc:mysql://([^:/]+)(?::(\d+))?/([^?]+)(?:\?(.*))?', url)
    if not m:
        sys.exit('DB_URL no reconocida.')
    host, puerto, base, opciones = m.group(1), int(m.group(2) or 3306), m.group(3), m.group(4) or ''
    ssl = {'ssl': True} if re.search(r'ssl-?mode=(REQUIRED|VERIFY)', opciones, re.I) else None
    print(f'Conectando a {host}:{puerto}/{base} como {usuario}')
    return pymysql.connect(host=host, port=puerto, user=usuario, password=clave_db, database=base, charset='utf8mb4', ssl=ssl, autocommit=False)


def pares(cur, sql):
    cur.execute(sql)
    return cur.fetchall()


class Catalogos:
    def __init__(self, cur):
        self.tipos_doc = {clave(n): i for i, n in pares(cur, 'SELECT id, nombre FROM tipo_documento')}
        self.paises = {clave(n): i for i, n in pares(cur, 'SELECT id, nombre FROM pais')}
        self.grupos = {clave(n): i for i, n in pares(cur, 'SELECT id, nombre FROM lista_grupo')}
        self.listas = {clave(n): (i, es) for i, n, es in pares(cur, 'SELECT id, nombre, es_pep FROM tipo_lista')}

    def buscar_lista(self, nombre_o_clave, regex=None):
        if regex:
            for k, (i, _) in self.listas.items():
                if re.search(regex, k):
                    return i
        return self.listas.get(clave(nombre_o_clave), (None,))[0]


# ----------------------------------------------------------------------------- informe y carga
def resumen(acum, existentes, nuevas_listas, db_manchas_dup):
    nuevas = [k for k in acum.ent if (clave(k[0]), k[1]) not in existentes]
    print('\n================ INFORME ================')
    print(f'filas leídas: {sum(acum.filas.values())}   ' + '   '.join(f'{k}: {v}' for k, v in acum.filas.items()))
    total_rech = sum(acum.rechazos.values())
    print(f'filas rechazadas (no se cargan): {total_rech}')
    for motivo, n in acum.rechazos.most_common():
        print(f'   - {motivo}: {n}   (filas de ejemplo: {", ".join(map(str, acum.ejemplos.get(motivo, [])))})')
    if acum.avisos:
        print('avisos (se cargan igual):')
        for motivo, n in acum.avisos.most_common():
            print(f'   - {motivo}: {n}')
    print(f'personas y empresas distintas en los archivos: {len(acum.ent)}   -> nuevas: {len(nuevas)}   ya existían: {len(acum.ent) - len(nuevas)}')
    print('   por tipo de documento: ' + ', '.join(f'{k} {v}' for k, v in acum.por_tipo_doc.most_common()))
    print(f'apariciones en listas: {len(acum.manchas)}   -> nuevas: {len(acum.manchas) - db_manchas_dup}   repetidas en la base: {db_manchas_dup}')
    print('   por lista:')
    for lista, n in acum.por_lista.most_common():
        etiqueta = LISTAS[lista][0] if lista in LISTAS else lista.replace('PLANTILLA:', '')
        marca = '  [LISTA NUEVA]' if lista in nuevas_listas else ''
        print(f'      {n:>8}  {etiqueta}{marca}')
    return len(nuevas)


def indices_necesarios(cur):
    cur.execute("SELECT COUNT(*) FROM information_schema.statistics WHERE table_schema = DATABASE() AND table_name = 'entidades' AND index_name = 'idx_entidades_documento'")
    if cur.fetchone()[0] == 0:
        print('   creando el índice por documento (hace rápida la búsqueda por DNI, RUC, etc.)...')
        cur.execute('CREATE INDEX idx_entidades_documento ON entidades (documento, id_tipo_documento)')
        return 1
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--archivo', action='append', required=True, help='Excel o CSV (se puede repetir)')
    ap.add_argument('--aplicar', action='store_true', help='escribe de verdad; sin esto solo simula')
    ap.add_argument('--env', default=None, help='archivo .env con DB_URL, DB_USER, DB_PASSWORD')
    ap.add_argument('--db-url'), ap.add_argument('--usuario'), ap.add_argument('--clave')
    ap.add_argument('--sin-base', action='store_true', help='solo lee y revisa los archivos, sin conectarse (para probar)')
    ap.add_argument('--rechazados', help='archivo CSV donde se guardan las filas rechazadas (contiene datos personales: guárdalo con cuidado)')
    ap.add_argument('--omitir-hoja', action='append', default=[], help='nombre de una hoja que no se carga (se puede repetir)')
    args = ap.parse_args()
    inicio = time.time()

    conexion = cur = cat = None
    if not args.sin_base:
        conexion = conectar(args)
        cur = conexion.cursor()
        cat = Catalogos(cur)

    f_rech = open(args.rechazados, 'w', encoding='utf-8-sig', newline='') if args.rechazados else None
    if f_rech:
        os.chmod(args.rechazados, 0o600)
    salida = csv.writer(f_rech) if f_rech else None
    if salida:
        salida.writerow(['archivo', 'hoja', 'fila', 'motivo', 'valores de la fila...'])

    acum = Acum()
    catalogo_plantilla = set(cat.listas) if cat else None
    omitir = set(args.omitir_hoja) | {'EJEMPLOS'}      # la hoja de ejemplos de la plantilla nunca se carga
    for ruta in args.archivo:
        procesar_archivo(ruta, acum, salida, catalogo_plantilla, omitir)
    if f_rech:
        f_rech.close()
    print(f'\nlectura terminada en {time.time() - inicio:.0f} s')

    # ----- que existe ya en la base
    existentes, db_manchas = {}, set()
    nuevas_listas = set()
    if cur:
        for i, t, d in pares(cur, 'SELECT e.id, td.nombre, e.documento FROM entidades e JOIN tipo_documento td ON td.id = e.id_tipo_documento'):
            existentes[(clave(t), d)] = i
        for r in pares(cur, 'SELECT id_entidades, id_tipo_lista, descripcion, cargo, institucion, fecha_registro FROM historial_manchas'):
            db_manchas.add(r)
        for lista in acum.por_lista:
            if lista in LISTAS and cat.buscar_lista(LISTAS[lista][0], LISTAS[lista][1]) is None:
                nuevas_listas.add(lista)
    dup = 0
    if cur and existentes:
        ids_lista = {}
        for lista in acum.por_lista:
            ids_lista[lista] = cat.buscar_lista(*(LISTAS[lista][:2] if lista in LISTAS else (lista.replace('PLANTILLA:', ''), None)))
        for k, lista, desc, _, f_desde, _, inst, cargo, *_ in acum.manchas:
            eid = existentes.get((clave(k[0]), k[1]))
            if eid and (eid, ids_lista.get(lista), desc, cargo, inst, f_desde) in db_manchas:
                dup += 1
    n_nuevas = resumen(acum, existentes, nuevas_listas, dup)

    if not args.aplicar:
        print('\nSIMULACIÓN: no se escribió nada. Si el informe está bien, repite con --aplicar.')
        return
    if sum(acum.filas.values()) == 0:
        sys.exit('No se leyó ninguna fila: no hay nada que cargar.')

    # ----- aplicar
    print('\n================ CARGA ================')
    t0 = time.time()
    # 1. catalogos que faltan
    for nombre, desc in TIPOS_DOC.items():
        if clave(nombre) not in cat.tipos_doc:
            cur.execute('INSERT INTO tipo_documento (nombre, descripcion) VALUES (%s, %s)', (nombre, desc))
            cat.tipos_doc[clave(nombre)] = cur.lastrowid
    for orden, (nombre, color) in GRUPOS.items():
        if clave(nombre) not in cat.grupos:
            cur.execute('INSERT INTO lista_grupo (orden, nombre, color) VALUES (%s, %s, %s)', (orden, nombre, color))
            cat.grupos[clave(nombre)] = cur.lastrowid
    ids_lista = {}
    for lista in acum.por_lista:
        if lista in LISTAS:
            nombre, regex, grupo, alcance, es_pep, desc = LISTAS[lista]
            i = cat.buscar_lista(nombre, regex)
            if i is None:
                cur.execute('INSERT INTO tipo_lista (nombre, descripcion, fuente, alcance, es_pep, id_grupo) VALUES (%s, %s, %s, %s, %s, %s)',
                            (nombre, desc, 'Carga masiva', alcance, es_pep, cat.grupos[clave(GRUPOS[grupo][0])]))
                i = cur.lastrowid
                print(f'   lista nueva creada: {nombre}')
        else:
            i = cat.buscar_lista(lista.replace('PLANTILLA:', ''))
        ids_lista[lista] = i
    sin_id = [k for k, v in ids_lista.items() if v is None]
    if sin_id:
        sys.exit('No se encontró en la base la lista: ' + ', '.join(sin_id))
    conexion.commit()
    cat_pais = dict(cat.paises)

    def id_pais(nombre):
        if not nombre:
            return None
        k = clave(nombre)
        if k not in cat_pais:
            cur.execute('INSERT INTO pais (nombre) VALUES (%s)', (nombre,))
            cat_pais[k] = cur.lastrowid
        return cat_pais[k]

    # 2. personas y empresas nuevas, por lotes
    ids_ent = {}
    nuevas = [(k, v) for k, v in acum.ent.items() if (clave(k[0]), k[1]) not in existentes]
    for k in acum.ent:
        e = existentes.get((clave(k[0]), k[1]))
        if e:
            ids_ent[k] = e
    tipo_por_id = {i: n for n, i in cat.tipos_doc.items()}
    for ini in range(0, len(nuevas), LOTE):
        lote = nuevas[ini:ini + LOTE]
        cur.execute('SELECT COALESCE(MAX(id), 0) FROM entidades')
        tope = cur.fetchone()[0]
        cur.executemany('INSERT INTO entidades (id_tipo_documento, id_pais, documento, tipo_entidad, fecha_registro, alias) VALUES (%s, %s, %s, %s, %s, %s)',
                        [(cat.tipos_doc[clave(k[0])], id_pais(v[1]), k[1], v[0], v[8], v[6]) for k, v in lote])
        cur.execute('SELECT id, id_tipo_documento, documento FROM entidades WHERE id > %s', (tope,))
        nuevos_ids = {(tipo_por_id[t], d): i for i, t, d in cur.fetchall()}
        naturales, juridicas = [], []
        for k, v in lote:
            eid = nuevos_ids[(clave(k[0]), k[1])]
            ids_ent[k] = eid
            if v[0] == 'NATURAL':
                naturales.append((eid, v[4], v[2], v[3], v[7]))
            else:
                juridicas.append((eid, v[5]))
        if naturales:
            cur.executemany('INSERT INTO personas_naturales (id_entidades, nombre, ape_pat, ape_mat, sexo) VALUES (%s, %s, %s, %s, %s)', naturales)
        if juridicas:
            cur.executemany('INSERT INTO personas_juridicas (id_entidades, razon_social) VALUES (%s, %s)', juridicas)
        conexion.commit()
        if (ini // LOTE) % 10 == 0 or ini + LOTE >= len(nuevas):
            print(f'   personas y empresas: {min(ini + LOTE, len(nuevas))} de {len(nuevas)}  ({time.time() - t0:.0f} s)')

    # 3. apariciones en listas, por lotes
    pendientes, escritas = [], 0
    for k, lista, desc, link, f_desde, f_hasta, inst, cargo, tipo_pep, p_desde, p_hasta in acum.manchas:
        eid, lid = ids_ent[k], ids_lista[lista]
        if (eid, lid, desc, cargo, inst, f_desde) in db_manchas:
            continue
        pendientes.append((eid, lid, desc, link, f_desde, f_hasta, inst, cargo, tipo_pep, p_desde, p_hasta))
        if len(pendientes) >= LOTE:
            cur.executemany('INSERT INTO historial_manchas (id_entidades, id_tipo_lista, descripcion, link, fecha_registro, fecha_hasta, institucion, cargo, tipo_pep, '
                            'periodo_desde, periodo_hasta) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)', pendientes)
            conexion.commit()
            escritas += len(pendientes)
            pendientes = []
            if (escritas // LOTE) % 10 == 0:
                print(f'   apariciones: {escritas}  ({time.time() - t0:.0f} s)')
    if pendientes:
        cur.executemany('INSERT INTO historial_manchas (id_entidades, id_tipo_lista, descripcion, link, fecha_registro, fecha_hasta, institucion, cargo, tipo_pep, '
                        'periodo_desde, periodo_hasta) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)', pendientes)
        conexion.commit()
        escritas += len(pendientes)
    print(f'   apariciones: {escritas}  ({time.time() - t0:.0f} s)')

    # 4. indice para que la busqueda por documento sea rapida
    indices_necesarios(cur)
    conexion.commit()
    for tabla in ('entidades', 'personas_naturales', 'personas_juridicas', 'historial_manchas'):
        cur.execute(f'ANALYZE TABLE {tabla}')
        cur.fetchall()
    print(f'\nAPLICADO: {len(nuevas)} personas/empresas nuevas y {escritas} apariciones en listas ({time.time() - inicio:.0f} s en total).')
    for tabla in ('entidades', 'personas_naturales', 'personas_juridicas', 'historial_manchas', 'tipo_lista'):
        cur.execute(f'SELECT COUNT(*) FROM {tabla}')
        print(f'   {tabla}: {cur.fetchone()[0]} filas')


if __name__ == '__main__':
    main()
